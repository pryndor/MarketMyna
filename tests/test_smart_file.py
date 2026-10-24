"""Smart messaging's report file: the same key figures every time, as PDF or
Excel, sent as an attachment / Telegram document. Admins and the owner get
calls; a general user gets readings of the indicators, never a call."""

from __future__ import annotations

import io
from datetime import datetime
from types import SimpleNamespace

import pytest

from marketmyna.services import smart_file
from marketmyna.services.report import ReportPrefs
from marketmyna.services.smart_file import build_document, report_file, to_pdf, to_xlsx
from marketmyna.services.smart_report import without_calls


def _row(symbol, action="buy", **extra):
    return {"symbol": symbol, "market": "india", "name": f"{symbol} Ltd", "price": 100.0, "change_percent": 1.5,
            "normal_low": 95.0, "normal_high": 105.0, "buy_point": 97.0, "sell_point": 104.0, "stop_loss": 94.0,
            "week_low": 97.0, "week_high": 103.0, "rsi": 55.0, "trend": "uptrend", "valuation": "fair", "pe": 21.0,
            "volatility": 0.22, "max_drawdown": -0.18, "high_52w": 120.0, "low_52w": 80.0,
            "action": action, "reasons": ["at its buy point"], **extra}


ROWS = [_row("TCS", ai=True, ai_summary="Strong order book."), _row("INFY", "hold"), _row("WIPRO", "sell"),
        {"symbol": "BAD", "market": "india", "error": "no history"}]
NOW = datetime(2026, 10, 17, 9, 0)


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    """No network: an index level, and valuation vs sector + events per stock."""
    monkeypatch.setattr("marketmyna.services.analysis_run.market_overview",
                        lambda markets: [{"name": "NIFTY 50", "price": 25000.0, "change_percent": 0.5}])
    monkeypatch.setattr("marketmyna.services.report.market_review", lambda overview, results, config: "Broad rise.")
    monkeypatch.setattr(smart_file, "_extras", lambda rows: {
        "TCS": {"valuation": {"position": 0.8, "zone": "below", "sector": "Technology"},
                "events": {"earnings_date": "2026-10-24"}}})


def _doc(rows=None, **prefs):
    """The file as run_smart builds it: every row through without_calls."""
    rows = [without_calls(r) for r in ROWS] if rows is None else rows
    return build_document(rows, ReportPrefs(**prefs), None, NOW, "MarketMyna · 4 stocks")


def test_same_core_set_whatever_is_ticked():
    doc = _doc(sections=["news", "alerts", "funds"])  # none of these go in the file
    titles = [b.title for b in doc.blocks]
    assert titles == ["Market overview", "What the indicators show", "Stocks at a glance", "How to read this report"]
    assert titles == [b.title for b in _doc(sections=["calls", "price"]).blocks]
    readings = doc.blocks[1]
    assert readings.rows[0][:2] == ["▲ Bullish", "TCS · AI"] and readings.rows[2][:2] == ["▼ Bearish", "WIPRO"]
    stocks = doc.blocks[2]
    assert stocks.columns == ["Stock", "Name", "Indicators", "Price", "Day", "Valuation vs sector", "What stands out"]
    tcs = dict(zip(stocks.columns, stocks.rows[0]))
    assert tcs["Indicators"] == "▲ Bullish · AI" and tcs["What stands out"] == "Strong order book."
    assert tcs["Valuation vs sector"] == "Priced 20% below the Technology median"
    infy = dict(zip(stocks.columns, stocks.rows[1]))
    assert infy["Valuation vs sector"] == "P/E 21.0; no sector median to compare"
    assert stocks.rows[3][1].startswith("no data")


def test_summary_tiles_and_stock_cards():
    doc = _doc()
    tiles = {t.label: t for t in doc.highlights}
    assert tiles["Stocks"].value == "4" and tiles["Bullish"].value == "1" and tiles["Bearish"].value == "1"
    card = doc.cards[0]
    facts = dict(card.facts)
    assert card.symbol == "TCS" and card.call == "▲ Bullish · AI" and card.price == "₹100.00"
    assert "Stop loss" not in facts and facts["In 52-week range"] == "50% of range"
    assert facts["Volatility (yearly)"] == "22%" and facts["Next results"] == "2026-10-24"
    assert card.notes == [[("AI summary of the figures", "Strong order book.", "accent"),
                           ("Valuation vs sector", "Priced 20% below the Technology median", "accent")]]
    assert len(doc.cards) == 3 and doc.notice  # the stock without data has no card


def test_general_users_get_readings_never_calls():
    rows = [without_calls(r) for r in ROWS]
    assert "action" not in rows[0] and "buy_point" not in rows[0] and "stop_loss" not in rows[0]
    assert rows[0]["reading"] == "positive" and rows[2]["reading"] == "negative"
    assert rows[0]["reasons"] == ["near its support level"] and rows[0]["valuation"] is None
    doc = _doc(rows)
    assert doc.notice.startswith("This report describes what the market data shows")
    titles = [b.title for b in doc.blocks]
    assert "Buy / Hold / Sell" not in titles and "What the indicators show" in titles
    stocks = next(b for b in doc.blocks if b.title == "Stocks at a glance")
    tcs = dict(zip(stocks.columns, stocks.rows[0]))
    assert tcs["Indicators"] == "▲ Bullish · AI" and "Call" not in stocks.columns
    assert tcs["Valuation vs sector"] == "Priced 20% below the Technology median"  # no "undervalued"
    facts = dict(doc.cards[0].facts)
    assert not {"Stop loss", "Buy at or below", "Sell at or above"} & set(facts)
    assert {t.label for t in doc.highlights} >= {"Bullish", "Neutral", "Bearish"}
    text = " ".join(str(v) for b in doc.blocks for row in b.rows for v in row)
    for word in ("BUY", "SELL", "Undervalued", "Overvalued", "Stop"):
        assert word not in text.replace("buy, sell or hold", "")


def test_pdf_and_excel_render():
    import fitz
    from openpyxl import load_workbook

    for doc in (_doc(),):
        pdf = to_pdf(doc)
        assert pdf.startswith(b"%PDF")
        pages = fitz.open(stream=pdf, filetype="pdf")
        first = pages[0].get_text()
        assert "WIPRO" in first and f"Page 1 of {len(pages)}" in first
        assert "₹100.00" in first  # the rupee sign comes from the fallback font
        assert any("Stock details" in page.get_text() for page in pages)
        assert ("PLEASE NOTE" in first) == bool(doc.notice)

    wb = load_workbook(io.BytesIO(to_xlsx(_doc())))
    assert wb.sheetnames == ["Summary", "Stocks", "How to read this report", "Stock details"]
    tcs = {h.value: c for h, c in zip(wb["Stocks"][1], wb["Stocks"][2])}
    assert tcs["Stock"].value == "TCS" and tcs["Price"].value == 100 and tcs["Day"].value == 0.015
    assert "0.0%" in tcs["Day"].number_format  # numbers sort as numbers
    assert any(c.value == "What the indicators show" for c in wb["Summary"]["A"])
    details = {h.value: c.value for h, c in zip(wb["Stock details"][1], wb["Stock details"][2])}
    assert details["AI summary of the figures"] == "Strong order book." and "Stop loss" not in details


def test_wide_tables_fit_the_page_without_breaking_numbers():
    from marketmyna.services.smart_pdf import CONTENT_W, Canvas

    columns = [f"Column {i}" for i in range(30)]
    cells = [["1,234,567.89"] * 30]
    numeric = [True] * 30
    assert Canvas._widths(columns, cells, 7.5, 4, numeric, set()) is None  # the caller shrinks the font
    widths = Canvas._widths(columns, cells, 5.5, 4, numeric, set(), force=True)
    assert abs(sum(widths) - CONTENT_W) < 1


def test_report_file_format_and_none():
    prefs = ReportPrefs(sections=["calls", "price"], smart_file="xlsx")
    name, data = report_file(ROWS, prefs, None, NOW, "t")
    assert name == "MarketMyna-report-2026-10-17.xlsx" and data[:2] == b"PK"
    assert report_file(ROWS, ReportPrefs(sections=["calls"], smart_file="none"), None, NOW, "t") is None


def test_telegram_sends_the_file_as_a_document(monkeypatch):
    from marketmyna.notifiers import telegram

    calls = []
    monkeypatch.setattr(telegram.requests, "post",
                        lambda url, **kw: calls.append((url, kw)) or SimpleNamespace(status_code=200, text=""))
    sender = telegram.TelegramSender(SimpleNamespace(telegram_bot_token="T", telegram_chat_id="1,2",
                                                     telegram_message_thread_id=None))
    assert sender.send_document("r.pdf", b"%PDF", caption="Report")
    assert [c[0].endswith("/sendDocument") for c in calls] == [True, True]
    assert calls[0][1]["files"]["document"] == ("r.pdf", b"%PDF") and calls[1][1]["data"]["chat_id"] == "2"


def test_email_attaches_excel_with_its_type(monkeypatch):
    from marketmyna.notifiers import email

    sent = []

    class FakeSMTP:
        def __init__(self, *a, **k):
            pass

        def login(self, *a):
            pass

        def sendmail(self, frm, to, msg):
            sent.append(msg)

        def send_message(self, msg, *a, **k):
            sent.append(msg.as_string())

        def quit(self):
            pass

    monkeypatch.setattr(email.smtplib, "SMTP_SSL", FakeSMTP)
    sender = email.EmailSender(SimpleNamespace(email_sender="me@example.com", email_login="me@example.com",
                                               email_password="x", email_receivers=["you@example.com"],
                                               email_smtp_server="smtp.example.com", email_smtp_port=465,
                                               email_smtp_ssl=True, email_sender_name=None))
    assert sender.send("body", title="Report", attachments=[("r.xlsx", b"PK..")])
    assert "spreadsheetml.sheet" in sent[0] and 'filename="r.xlsx"' in sent[0]
