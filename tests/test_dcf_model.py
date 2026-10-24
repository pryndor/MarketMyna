"""Tests for the DCF Excel model (structure; values were checked against
services/dcf.model_value with an independent formula engine)."""

import io

import pytest
from openpyxl import load_workbook

from src.services import dcf_model

DATA = {"symbol": "ACME", "market": "india", "name": "Acme Ltd", "currency": "INR", "price": 100.0,
        "periods": ["2026-03-31", "2025-03-31"], "source": "Test filings",
        "dcf": {"base_fcf": 1000.0, "base_label": "latest year", "net_cash": 200.0, "shares": 50.0,
                "fcf_history": [1000.0, 900.0], "converted_from": None,
                "defaults": {"growth": 0.08, "discount": 0.12, "terminal": 0.05, "high_years": 5, "total_years": 10}}}


def test_inputs_formulas_and_checks():
    wb = load_workbook(io.BytesIO(dcf_model.build(DATA, discount=0.11)))
    ws = wb["DCF"]
    assert [ws[f"B{r}"].value for r in range(6, 13)] == [1000.0, 0.08, 0.11, 0.05, 200.0, 50.0, 100.0]
    assert ws["B6"].font.color.rgb.endswith("1F4FBF")  # inputs are blue
    assert ws["C15"].value == "=$B$7" and ws["H15"].value.startswith("=$B$7+($B$9-$B$7)")
    assert ws["D16"].value == "=C16*(1+D15)"
    assert ws["B26"].value == "=B25/B11" and "OK" in ws["B29"].value
    # every calculation cell is a formula, never a typed number
    for row in list(range(15, 19)) + list(range(33, 38)):
        for col in "CDEFG":
            assert str(ws[f"{col}{row}"].value).startswith("="), f"{col}{row}"
    assert "SUMPRODUCT" in ws["E35"].value and '"n/a"' in ws["E35"].value
    assert [wb["History"][f"B{r}"].value for r in (2, 3)] == [1000.0, 900.0]


def test_no_dcf_raises():
    with pytest.raises(ValueError, match="banks"):
        dcf_model.build({"symbol": "BANK", "dcf": None, "note": "No DCF for banks and insurers"})
